"""Structured Word (.docx) and PowerPoint (.pptx) reports of a crawl job."""
from __future__ import annotations

import io
import time
from typing import TYPE_CHECKING

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptColor
from pptx.util import Emu, Inches
from pptx.util import Pt as PPt

if TYPE_CHECKING:
    from .jobs import Job

FONT = "Microsoft YaHei"
ACCENT = (0x2F, 0x6F, 0xED)
SOURCE_LABEL = "来源页面"
STATUS = {"queued": "排队中", "running": "进行中", "done": "已完成", "failed": "失败", "cancelled": "已取消"}


def columns(job: Job) -> list[tuple[str, str]]:
    """(key, label) for every column: planned fields, unexpected extras, then the source URL."""
    fields = (job.plan or {}).get("fields", [])
    cols = [(f["key"], f["label"]) for f in fields]
    known = {k for k, _ in cols}
    extra = sorted({k for r in job.records for k in r} - known - {"_source"})
    return cols + [(k, k) for k in extra] + [("_source", SOURCE_LABEL)]


def _fmt_time(ts: float | None) -> str:
    return time.strftime("%Y-%m-%d %H:%M", time.localtime(ts)) if ts else "-"


def _clip(text: str, n: int) -> str:
    text = " ".join(str(text or "").split())
    return text if len(text) <= n else text[: n - 1] + "…"


def _use_table(job: Job, cols: list[tuple[str, str]]) -> bool:
    """Short, few-column records read best as a table; any long column -> one section per record."""
    data_cols = [k for k, _ in cols if k != "_source"]
    if len(data_cols) > 6:
        return False
    if not job.records:
        return True
    longest = max(sum(len(str(r.get(k, ""))) for r in job.records) / len(job.records) for k in data_cols)
    return longest <= 80


def _overview(job: Job) -> list[tuple[str, str]]:
    plan = job.plan or {}
    rows = [
        ("抓取需求", job.request.description or "（未填写）"),
        ("起始网址", "\n".join(job.request.urls)),
    ]
    if plan:
        rows += [("AI 理解", plan.get("summary", "")), ("每条记录", plan.get("item_description", ""))]
    ok = sum(1 for p in job.pages if p.get("status") == "ok")
    rows += [
        ("任务状态", STATUS.get(job.status, job.status)),
        ("抓取页面", f"{len(job.pages)} 个（成功 {ok} 个）"),
        ("记录数量", f"{len(job.records)} 条"),
        ("创建时间", _fmt_time(job.created_at)),
        ("完成时间", _fmt_time(job.finished_at)),
    ]
    return rows


# ---------------------------------------------------------------- Word

def _docx_font(run, size: float | None = None, bold: bool = False, color=None) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def _docx_heading(doc, text: str, level: int) -> None:
    h = doc.add_heading(level=level)
    _docx_font(h.add_run(text), color=ACCENT if level == 1 else None)


def _docx_para(doc_or_cell, text: str, size: float = 10.5, bold: bool = False, color=None):
    p = doc_or_cell.add_paragraph()
    _docx_font(p.add_run(text), size, bold, color)
    return p


def _set_cell(cell, text: str, bold: bool = False, size: float = 9) -> None:
    cell.text = ""
    _docx_font(cell.paragraphs[0].add_run(text), size, bold)


def _shade(cell, hex_fill: str) -> None:
    from docx.oxml import OxmlElement

    tc_pr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def to_docx(job: Job) -> bytes:
    cols = columns(job)
    as_table = _use_table(job, cols)
    doc = Document()
    sec = doc.sections[0]
    if as_table and len(cols) > 4:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2))

    title = doc.add_heading(level=0)
    _docx_font(title.add_run(_clip(job.request.description or "网页抓取报告", 60)))
    _docx_para(doc, f"Craw Pro 抓取报告 · 生成于 {_fmt_time(time.time())}", 9, color=(0x6B, 0x74, 0x82))

    # 1. Overview
    _docx_heading(doc, "一、任务概览", 1)
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in _overview(job):
        c1, c2 = t.add_row().cells
        _set_cell(c1, k, bold=True, size=10)
        _shade(c1, "E8F0FE")
        _set_cell(c2, v, size=10)
        c1.width, c2.width = Cm(3.5), Cm(13)

    # 2. Fields
    fields = (job.plan or {}).get("fields", [])
    if fields:
        _docx_heading(doc, "二、提取字段", 1)
        t = doc.add_table(rows=1, cols=2)
        t.style = "Table Grid"
        for cell, txt in zip(t.rows[0].cells, ("字段", "说明")):
            _set_cell(cell, txt, bold=True, size=10)
            _shade(cell, "E8F0FE")
        for f in fields:
            c1, c2 = t.add_row().cells
            _set_cell(c1, f["label"], size=10)
            _set_cell(c2, f.get("description", ""), size=10)

    # 3. Records
    _docx_heading(doc, f"三、抓取结果（共 {len(job.records)} 条）", 1)
    if not job.records:
        _docx_para(doc, "没有抓取到记录。")
    elif as_table:
        t = doc.add_table(rows=1, cols=len(cols) + 1)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = t.rows[0].cells
        for cell, txt in zip(hdr, ["#"] + [label for _, label in cols]):
            _set_cell(cell, txt, bold=True)
            _shade(cell, "E8F0FE")
        for i, r in enumerate(job.records, 1):
            cells = t.add_row().cells
            _set_cell(cells[0], str(i))
            for cell, (k, _) in zip(cells[1:], cols):
                _set_cell(cell, str(r.get(k, "")))
    else:
        first_key, first_label = cols[0]
        for i, r in enumerate(job.records, 1):
            _docx_heading(doc, f"{i}. {_clip(r.get(first_key, '') or f'记录 {i}', 80)}", 2)
            for k, label in cols:
                val = str(r.get(k, "") or "")
                if not val or (k == first_key and len(val) <= 80):
                    continue
                p = doc.add_paragraph()
                _docx_font(p.add_run(f"{label}："), 10.5, bold=True)
                _docx_font(p.add_run(val), 10.5)

    # 4. Pages
    if job.pages:
        _docx_heading(doc, "四、抓取页面", 1)
        t = doc.add_table(rows=1, cols=4)
        t.style = "Table Grid"
        for cell, txt in zip(t.rows[0].cells, ("页面", "标题", "状态", "记录数")):
            _set_cell(cell, txt, bold=True)
            _shade(cell, "E8F0FE")
        label = {"ok": "成功", "error": "失败", "blocked": "robots 禁止"}
        for p in job.pages:
            cells = t.add_row().cells
            _set_cell(cells[0], p.get("url", ""))
            _set_cell(cells[1], p.get("title", "") or "")
            _set_cell(cells[2], label.get(p.get("status"), p.get("status", "")))
            _set_cell(cells[3], str(p.get("records", 0)))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PowerPoint

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.5)


def _ppt_text(tf, text: str, size: float, bold: bool = False, color=None, first: bool = True):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    run = p.add_run()
    run.text = text
    run.font.size = PPt(size)
    run.font.bold = bold
    run.font.name = FONT
    if color:
        run.font.color.rgb = PptColor(*color)
    return p


def _ppt_slide(prs, title: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    bar = slide.shapes.add_shape(1, 0, 0, SLIDE_W, Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = PptColor(*ACCENT)
    bar.line.fill.background()
    tb = slide.shapes.add_textbox(MARGIN, Inches(0.3), SLIDE_W - 2 * MARGIN, Inches(0.8))
    tb.text_frame.word_wrap = True
    _ppt_text(tb.text_frame, title, 26, bold=True, color=(0x1C, 0x23, 0x30))
    return slide


def _ppt_table(slide, rows: list[list[str]], col_widths: list[float], top=Inches(1.3), font=11):
    height = SLIDE_H - top - Inches(0.4)
    shape = slide.shapes.add_table(len(rows), len(rows[0]), MARGIN, top, SLIDE_W - 2 * MARGIN, Emu(int(height)))
    table = shape.table
    total = sum(col_widths)
    for i, w in enumerate(col_widths):
        table.columns[i].width = Emu(int((SLIDE_W - 2 * MARGIN) * w / total))
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            _ppt_text(cell.text_frame, val, font, bold=(r == 0),
                      color=(0xFF, 0xFF, 0xFF) if r == 0 else None)
            cell.text_frame.word_wrap = True
            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = PptColor(*ACCENT)
    return table


def to_pptx(job: Job, max_record_slides: int = 200) -> bytes:
    cols = columns(job)
    data_cols = [c for c in cols if c[0] != "_source"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H

    # Title slide
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(1, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = PptColor(*ACCENT)
    bg.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(0.9), Inches(2.4), SLIDE_W - Inches(1.8), Inches(2.5))
    tb.text_frame.word_wrap = True
    _ppt_text(tb.text_frame, _clip(job.request.description or "网页抓取报告", 70), 36, True, (0xFF, 0xFF, 0xFF))
    _ppt_text(tb.text_frame, f"{len(job.records)} 条记录 · {len(job.pages)} 个页面 · {_fmt_time(job.finished_at or job.created_at)}",
              16, color=(0xE8, 0xF0, 0xFE), first=False)
    _ppt_text(tb.text_frame, _clip(" · ".join(job.request.urls), 120), 14, color=(0xE8, 0xF0, 0xFE), first=False)

    # Overview
    slide = _ppt_slide(prs, "任务概览")
    _ppt_table(slide, [["项目", "内容"]] + [[k, _clip(v, 200)] for k, v in _overview(job)], [1, 4], font=13)

    # Fields
    fields = (job.plan or {}).get("fields", [])
    if fields:
        slide = _ppt_slide(prs, "提取字段")
        rows = [["字段", "说明"]] + [[f["label"], _clip(f.get("description", ""), 120)] for f in fields[:12]]
        _ppt_table(slide, rows, [1, 3], font=13)

    # Records
    if not job.records:
        slide = _ppt_slide(prs, "抓取结果")
        tb = slide.shapes.add_textbox(MARGIN, Inches(1.5), SLIDE_W - 2 * MARGIN, Inches(1))
        _ppt_text(tb.text_frame, "没有抓取到记录。", 18)
    elif _use_table(job, cols):
        per = 8
        show = data_cols
        total = len(job.records)
        pages = (total + per - 1) // per
        for n in range(min(pages, max_record_slides)):
            chunk = job.records[n * per : (n + 1) * per]
            slide = _ppt_slide(prs, f"抓取结果（{n * per + 1}–{n * per + len(chunk)} / {total}）")
            rows = [["#"] + [label for _, label in show]]
            rows += [[str(n * per + i + 1)] + [_clip(r.get(k, ""), 80) for k, _ in show] for i, r in enumerate(chunk)]
            _ppt_table(slide, rows, [0.4] + [2] * len(show), font=11)
    else:
        first_key = data_cols[0][0] if data_cols else "_source"
        for i, r in enumerate(job.records[:max_record_slides], 1):
            slide = _ppt_slide(prs, f"{i}. {_clip(r.get(first_key, '') or f'记录 {i}', 50)}")
            tb = slide.shapes.add_textbox(MARGIN, Inches(1.3), SLIDE_W - 2 * MARGIN, SLIDE_H - Inches(1.7))
            tf = tb.text_frame
            tf.word_wrap = True
            first = True
            for k, label in cols:
                val = str(r.get(k, "") or "")
                if not val:
                    continue
                p = _ppt_text(tf, f"{label}：", 14, bold=True, color=ACCENT, first=first)
                run = p.add_run()
                run.text = _clip(val, 400)
                run.font.size = PPt(14)
                run.font.name = FONT
                p.space_after = PPt(6)
                first = False

    shown = min(len(job.records), max_record_slides * (8 if _use_table(job, cols) else 1))
    slide = _ppt_slide(prs, "说明")
    tb = slide.shapes.add_textbox(MARGIN, Inches(1.4), SLIDE_W - 2 * MARGIN, Inches(3))
    tb.text_frame.word_wrap = True
    _ppt_text(tb.text_frame, f"本演示文稿展示了 {shown} / {len(job.records)} 条记录，较长的内容已截断。", 18)
    _ppt_text(tb.text_frame, "完整数据请在 Craw Pro 中导出 Word、CSV 或 JSON。", 18, first=False)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
